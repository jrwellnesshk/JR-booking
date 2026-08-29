# -*- coding: utf-8 -*-
"""生成場景測試報告 Word 檔"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DARK = RGBColor(0x1F, 0x3B, 0x57)
RED = RGBColor(0xC0, 0x22, 0x1D)
AMBER = RGBColor(0xB4, 0x6A, 0x00)
GREEN = RGBColor(0x1E, 0x7A, 0x33)
GREY = RGBColor(0x66, 0x66, 0x66)

doc = Document()

# 全域字體
style = doc.styles['Normal']
style.font.name = 'Microsoft JhengHei'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')

for sec in doc.sections:
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)

def set_cn(run):
    run.font.name = 'Microsoft JhengHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
    return run

def para(text='', size=10.5, bold=False, color=None, space_after=6, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align: p.alignment = align
    if text:
        r = set_cn(p.add_run(text))
        r.font.size = Pt(size); r.bold = bold
        if color: r.font.color.rgb = color
    return p

def heading(text, level=1):
    sizes = {1: 16, 2: 13, 3: 11.5}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = set_cn(p.add_run(text))
    r.bold = True; r.font.size = Pt(sizes.get(level, 11)); r.font.color.rgb = DARK
    return p

def bullet(text, bold_prefix=None, color=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    if bold_prefix:
        r = set_cn(p.add_run(bold_prefix)); r.bold = True; r.font.size = Pt(10.5)
        if color: r.font.color.rgb = color
    r2 = set_cn(p.add_run(text)); r2.font.size = Pt(10.5)
    return p

def mono(text, color=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(text)
    r.font.name = 'Consolas'; r.font.size = Pt(9)
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    if color: r.font.color.rgb = color
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        r = set_cn(hdr[i].paragraphs[0].add_run(h)); r.bold = True; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = set_cn(cells[i].paragraphs[0].add_run(str(v))); r.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t

# ================= 封面標題 =================
para('寶天醫館預約系統（booking-aurora）', 12, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para('30 客全流程場景測試報告', 24, bold=True, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
para('測試日期：2026年8月24日　|　環境：localhost:4000　|　模擬規模：一日30位客人（對應每月約400客之高峯日）', 9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para('測試方式：自動化 E2E API 測試（237 個調用 / 50 秒）　|　總結果：180 通過 / 1 項例外（證實為系統 bug）', 9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

# ================= 一、場景設計 =================
heading('一、30 客場景設計')
table(
    ['組別', '人數', '模擬情境'],
    [
        ['訪客 V1–V3', '3', '無帳戶網上預約初體驗（S1）'],
        ['註冊新客 zta01–07', '7', '網上開戶→補資料→預約；含改密碼、自助取消、寫評價/論壇'],
        ['舊客 zc01–04＋8udc', '5', '職員建立舊檔；premium/family 升級；12歲子帳戶全流程'],
        ['散客 zd01–10', '10', '職員現場代訂；完成×4／pending／no-show／職員取消／遲到15分／配藥中'],
        ['高級會員 ze01–05', '5', '床位S6、小兒推拿(8歲)、自行改期、取消窗口邊緣測試'],
    ],
    widths=[4, 1.5, 11]
)
para()
para('最終數據庫狀態（API 回報與 SQLite 實際記錄一致）：32 個測試預約＝confirmed 19／completed 7／cancelled 5／no-show 1；26 個測試用戶；8 條會員訂閱；1 個家庭連結；2 份病歷。', 10)

# ================= 二、通過功能 =================
heading('二、通過驗證嘅功能（無需修改）')
groups = [
    ('預約核心', ['同醫師同時段雙訂 → 409 正確拒絕', '過去日期、週日休診、非營業時間全部攔截', '床位服務（S6）容量邏輯正常；時段可用性與實際一致']),
    ('會員分級', ['訪客只可約 S1；一般會員升級前只可約 S1（403 membership_required）', 'Stripe 未配置時 checkout 正確回錯誤']),
    ('帳戶安全', ['重複用戶名／電話／弱密碼／無中文名註冊全被擋', '登出後 token 即時失效；角色越權全部 403', '敏感操作二次驗證生效']),
    ('診所運作', ['狀態機流轉：pending→confirmed→in-treatment→completed、visited→dispensing、no-show', '醫師請假→當日時段封鎖→取消即恢復', '家庭子帳戶：開戶(<18)→臨時密碼登入→重設→解除→重連，全鏈路通']),
    ('臨床記錄', ['病歷建立→進度指標→AI 分析→客人自查病史→自助記錄感受', '缺必填欄位有明確驗證提示']),
    ('互動內容', ['回饋閉環：提交→回覆→未讀數→已讀→結案', '分診/AI 問診答題→建議結果', '評價審核後公開；論壇發帖+回覆；FAQ/公告 CRUD']),
    ('報表', ['日/月收入、初體驗清單、職員當日總覽、密碼重設日誌、可疑活動檢查']),
]
for g, items in groups:
    heading(g, 3)
    for it in items:
        bullet(it)

# ================= 三、問題 =================
heading('三、發現嘅問題（按嚴重性排序）')

heading('🔴 高優先', 2)
p = para('1. 通知系統「設定」與「實際行為」不符（有金錢風險）', 11.5, bold=True, color=RED, space_after=3)
bullet('WhatsApp 行緊 Twilio Sandbox 模式：每個預約確認都發送失敗——真客而家收唔到任何 WhatsApp 確認（server.log 證據）。', bold_prefix='')
bullet('WhatsApp 失敗後 fallback 去 ClickSend 發真實 SMS——雖然 README 寫「SMS 已全面取消」、通知設定 SMS=false，日誌仍顯示每個測試預約都提交咗 ClickSend SMS。')
bullet('建議：① 即刻檢查 ClickSend 帳戶有冇被扣費；② 移除 SMS fallback 或令 fallback 尊重 sms_notification_enabled 開關；③ 用正式 Twilio WhatsApp 號碼取代 Sandbox。（按負責人指示，通知功能未做深入測試，以上為日誌自然暴露證據）')

p = para('2. database.db 放喺 OneDrive 同步資料夾＋兩邊同時作業', 11.5, bold=True, color=RED, space_after=3)
bullet('實測發生：測試中途更改嘅 admin 密碼，幾分鐘後被另一邊同步覆寫返舊值。')
bullet('SQLite 係單一檔案，雲同步係成隻檔案替換——兩邊都有寫入嘅話，隨時成批新預約/會員資料被舊版本蓋走，而且唔會有任何報錯。')
bullet('建議：正式營運將資料庫移出 OneDrive 目錄（例如 C:\\clinic-data\\），只同步備份副本唔同步活體檔案；或規定只有一部機可以寫入。')

p = para('3. 客戶取消時間窗存在時區 bug — routes/bookings.js:884', 11.5, bold=True, color=RED, space_after=3)
mono("const created = new Date(booking.created_at + (booking.created_at.includes('T') ? '' : 'Z'));")
bullet('created_at 存本地時間字串，但無 T 就加 Z 尾當 UTC 計。香港（UTC+8）伺服器上「距離建立幾耐」會計多咗 8 小時——建立後頭約 8 小時內結果係負數，15 分鐘冷靜期檢查形同虛設，客人開單後差唔多 8 小時內可以取消任何 >24h 遠期預約。')
bullet('專項驗證：新單即刻取消 ✅200；created_at 回撥9小時 → ✅400 cancel_window 正確阻擋；只回撥2小時（<8h偏移）→ ❌200 取消成功（bug 重現）。')
mono('修法：new Date(booking.created_at.replace(\' \', \'T\'))   // 按本地時間解析，不加 Z', GREEN)
bullet('另請留意：而家規則係「預約前 24 小時內先可以取消」（代碼同錯誤訊息一致），即鼓勵臨急取消、反而禁止早啲通知嘅取消——一般診所做法相反（最少 24 小時前通知）。請確認業務意圖。')

heading('🟡 中優先', 2)
p = para('4. 職員代訂被「職員自己」嘅會員等級卡住', 11.5, bold=True, color=AMBER, space_after=3)
bullet('POST /api/bookings 用 token 持有人嘅 tier 判斷服務檔次：一般級 staff01 代客訂 S2–S6 → 403，錯誤訊息仲叫「客人自己去 Stripe 升級」——櫃檯場景完全行唔通。')
bullet('同時預約 user_id 記錄為建立者（職員）而非客人，報表/收入/病歷歸屬失真。建議：staff/admin 跳過 tier 檢查；支援 body.customerUserId（限 staff+）。')

p = para('5. /api/membership/confirm 嘅 userId fallback 危險', 11.5, bold=True, color=AMBER, space_after=3)
bullet('body.userId 缺失時自動升級「當前登入者」。測試期間曾意外將 staff01 升級做 premium（已還原）。建議 userId 必填。')

p = para('6. 登入速率限制對櫃檯偏緊', 11.5, bold=True, color=AMBER, space_after=3)
bullet('20 次/15分鐘/IP。測試兩度觸發 429。月 400 宵櫃檯幾個職員共用 IP 好易撞牆。建議改「帳戶+IP」組合，或 staff/doctor 用較高上限。')

p = para('7. FAQ API 路徑三套並存', 11.5, bold=True, color=AMBER, space_after=3)
bullet('公開讀取 /api/faqs（misc）、內容組 /api/content/*、管理 /api/admin/faqs——前端易出錯，建議統一。')

heading('🟢 低優先 / 觀察', 2)
for t in [
    '子帳戶 family_head_id 不一致：/family/register 有寫、/family/add 冇寫（功能靠 links 表所以暫時冇壞）。',
    '子帳戶臨時密碼唔強制首登更改（must_change_password 只係提示）。',
    '電話格式校驗寬鬆：任何 8 位數字都收，可加 HK 號碼範圍（6/9 開頭）。',
    '全域限制 500 req/15min/IP：今次 237 calls 冇事，前端若加輪詢要留意。',
]:
    bullet(t)

# ================= 四、建議 =================
heading('四、功能增強建議（配合每月 400 客規模）')
table(
    ['#', '建議', '價值', '難度'],
    [
        ['1', '候補名單（waitlist）：滿格時段登記候補，有人取消自動 WhatsApp 通知', '高峯日提升填單率', '中'],
        ['2', '訊息發送狀態面板：admin 後台顯示每條訊息提交/送達/失敗＋一鍵重發', '針對高優先第1項，「以為通知咗其實冇」', '低'],
        ['3', '營運儀表板：新客vs回頭客比例、醫師佔比、熱門時段熱力圖、營收曲線', '數據已有欠呈現；400客/月好需要', '中'],
        ['4', '客人有限度自助改期：確認後、預約前24h 內准改一次', '減輕櫃檯電話量', '低'],
        ['5', '報表匯出 Excel/PDF：月尾交數俾股東/會計', 'xlsx 套件已有，成本低', '低'],
        ['6', '自動備份排程：每日複製 database.db 到另一磁碟＋保留14份', '配合高優先第2項', '極低'],
        ['7', '提醒發送日摘要：每晚彙總發送/失敗統計電郵管理員', '提早發現渠道故障', '低'],
        ['8', '敏感操作審計日誌頁面：刪用戶/密碼重設集中可查', '多員工營運必要', '低'],
        ['9', '平板優化 staff.html：大按鈕模式（報到/遲到/完成一撳即轉）', '現場效率', '中'],
    ],
    widths=[0.9, 8.5, 4.8, 1.5]
)

# ================= 五、未覆蓋 =================
heading('五、本次未覆蓋範圍（誠實申報）')
for t in [
    '通知批量任務（run-all／天氣／節氣推送）— 按負責人指示跳過',
    'Stripe 真實付款流程（金鑰未配置）',
    '病歷音檔/相片 multipart 上傳、頭像上傳、影片上傳、Excel 入賬導入',
    '純前端 UI 互動（今次集中 API 層）',
]:
    bullet(t)

# ================= 六、產物 =================
heading('六、測試產物與清理狀態')
table(
    ['檔案', '用途'],
    [
        ['_e2e_scenario_test.js', '主場景測試腳本（可重跑）'],
        ['_e2e_results.json', '完整逐步結果（181 條記錄）'],
        ['_tmp_cancel_test.js', '取消窗口 bug 專項驗證腳本'],
        ['_tmp_clean.js', '清理所有【場景測試】數據＋還原 staff01'],
        ['場景測試報告_2026-08-24.md / .docx', '本報告'],
    ],
    widths=[6.5, 10]
)
para()
para('清理狀態：26 個測試用戶、32 個測試預約及相關病歷/回饋/評價已全部清除；staff01 已還原 general；allow_public_registration 已還原 false；D0–D2 只剩 5 個真實預約。所有測試數據均帶【場景測試】標記，與真實客戶完全隔離。', 10)

doc.save('場景測試報告_2026-08-24.docx')
print('saved: 場景測試報告_2026-08-24.docx')
