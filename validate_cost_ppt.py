# -*- coding: utf-8 -*-
import glob, os
from pptx import Presentation

bases = [os.path.dirname(os.path.abspath(__file__)),
         r'C:\Users\Damian\Desktop\booking-demo - vs code（21-01-26)']
cands = []
for b in bases:
    for c in glob.glob(os.path.join(b, '**', '*.pptx'), recursive=True):
        if '營運' in os.path.basename(c):
            cands.append(c)
cands = [c for c in cands if 'presentation' not in c.replace('\\', '/').lower()]
assert cands, 'no cost pptx found: ' + repr(cands)
path = sorted(cands, key=os.path.getmtime)[-1]
p = Presentation(path)
print('file:', path)
print('slides:', len(p.slides._sldIdLst))
for i, s in enumerate(p.slides, 1):
    firsts = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            firsts.append(sh.text_frame.text.strip().split('\n')[0])
    print(i, '|', ' / '.join(firsts[:3]))
tbl = sum(1 for s in p.slides for sh in s.shapes if sh.has_table)
print('tables:', tbl)
empty = [i for i, s in enumerate(p.slides, 1)
         if not any(sh.has_text_frame and sh.text_frame.text.strip() for sh in s.shapes)]
print('empty slides:', empty)
